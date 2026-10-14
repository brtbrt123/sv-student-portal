"""Builds docs/Team-Monitoring-Sheet.xlsx: a template to track tasks and hours per member and to record peer evaluation.
Needs: pip install openpyxl. Run: python scripts/pdf/build_monitoring_sheet.py"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / 'docs' / 'Team-Monitoring-Sheet.xlsx'

TEAM = [
    ('Abarquez, Justinbril', 'Project Manager / Systems Analyst'),
    ('Cruzante, Aaron', 'Backend / Database Developer'),
    ('Canayong, Ron', 'Backend / Database Developer'),
    ('Esgana, Justine', 'Frontend Developer'),
    ('Espedilon, Arvien', 'Frontend Developer'),
    ('Dominado, Byrone', 'Frontend Developer'),
]
NAMES = [n for n, _ in TEAM]
STATUS = ['Not started', 'In progress', 'Done', 'Blocked']
PRIORITY = ['High', 'Medium', 'Low']

DARK, MID, LIGHT, WHITE = '3F3F46', 'D4D4D8', 'F4F4F5', 'FFFFFF'
thin = Side(style='thin', color='A1A1AA')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(bold=True, color=WHITE, name='Calibri', size=11)
HEAD_FILL = PatternFill('solid', fgColor=DARK)
SUB_FILL = PatternFill('solid', fgColor=MID)
BAND = PatternFill('solid', fgColor=LIGHT)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, CENTER, BORDER
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = Font(bold=True, size=16, name='Calibri')
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(italic=True, color='52525B', name='Calibri')


def body(ws, first, last, ncols, banded=True):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if cell.alignment is None or not cell.alignment.wrap_text:
                cell.alignment = WRAP
            if banded and (r - first) % 2 == 1:
                cell.fill = BAND


# ------------------------------------------------------------------------------------------------ Instructions
ws = wb.active
ws.title = 'Instructions'
title(ws, 'Team Monitoring Sheet', 'Centralized Student Portal and Automated Clearance System | Pamantasan ng Cabuyao | BS Computer Science')
rows = [
    ('How to use this workbook', ''),
    ('1. Team', 'Check the names and roles on the Team sheet. The lists on the other sheets read from it.'),
    ('2. Task Tracker', 'One row per task or deliverable. Pick the member and the status from the lists, enter the hours spent, and add the link to the pull request, issue or file. Update the status as work moves.'),
    ('3. Weekly Progress', 'At the end of each week or sprint, each member writes what was planned, what was completed, the hours, blockers and next steps.'),
    ('4. Meeting Log', 'Record each team meeting, the decisions and the action items with an owner and a due date.'),
    ('5. Peer Evaluation', 'Each member rates every teammate from 1 to 5 on the six criteria. The average per row is calculated. Ratings are given at the end of the project (or each sprint).'),
    ('6. Summary', 'Calculated automatically: tasks done, hours, average peer score and share of the work per member. Do not type in the grey cells.'),
    ('', ''),
    ('Rating scale (Peer Evaluation)', '5 = Excellent   4 = Good   3 = Satisfactory   2 = Needs improvement   1 = Poor'),
    ('Status values', 'Not started, In progress, Done, Blocked'),
    ('Tip', 'Use the filter arrows in the header row to see the tasks of one member or one sprint.'),
]
for i, (a, b) in enumerate(rows, 4):
    ws.cell(row=i, column=1, value=a).font = Font(bold=True, name='Calibri')
    ws.cell(row=i, column=2, value=b).alignment = WRAP
ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 110

# ------------------------------------------------------------------------------------------------ Team
wt = wb.create_sheet('Team')
title(wt, 'Team', 'Names feed the drop-down lists on the other sheets.')
header(wt, 4, ['#', 'Member', 'Role', 'Student number', 'Email', 'Main responsibility'], [5, 28, 36, 18, 32, 60])
for i, (n, r) in enumerate(TEAM, 1):
    wt.cell(row=4 + i, column=1, value=i)
    wt.cell(row=4 + i, column=2, value=n)
    wt.cell(row=4 + i, column=3, value=r)
body(wt, 5, 4 + len(TEAM), 6)
resp = {1: 'Backlog, sprint plan, system models, SDLC and test plan, documentation, releases',
        2: 'Database, RLS policies, triggers, authentication, security', 3: 'Database, RLS policies, triggers, authentication, security',
        4: 'UI/UX wireframes and module screens', 5: 'UI/UX wireframes and module screens', 6: 'UI/UX wireframes and module screens'}
for i, t in resp.items():
    wt.cell(row=4 + i, column=6, value=t)
TEAM_RANGE = f"Team!$B$5:$B${4 + len(TEAM)}"

def dv_member_new():
    return DataValidation(type='list', formula1=f'={TEAM_RANGE}', allow_blank=True)



dv_status = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True)
dv_priority = DataValidation(type='list', formula1='"' + ','.join(PRIORITY) + '"', allow_blank=True)
dv_rating = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True,
                           error='Enter a whole number from 1 to 5', errorTitle='Rating')

# ------------------------------------------------------------------------------------------------ Task tracker
wk = wb.create_sheet('Task Tracker')
title(wk, 'Task Tracker', 'One row per task. The first three rows are examples; replace or delete them.')
cols = ['Task ID', 'Sprint', 'Task / deliverable', 'Assigned to', 'Priority', 'Status', 'Start date', 'Due date', 'Date done', 'Hours spent', 'Link (PR, issue, file)', 'Remarks']
header(wk, 4, cols, [9, 9, 46, 24, 11, 14, 12, 12, 12, 11, 36, 40])
examples = [
    ('T-001', 'S1', 'Example: ERD and 3NF schema', NAMES[1], 'High', 'Done', '2026-10-01', '2026-10-02', '2026-10-02', 12, 'issue #', 'Example row'),
    ('T-002', 'S1', 'Example: wireframes of the Student screens', NAMES[3], 'Medium', 'Done', '2026-10-01', '2026-10-02', '2026-10-02', 8, 'Figma link', 'Example row'),
    ('T-003', 'S2', 'Example: login and role routing', NAMES[2], 'High', 'In progress', '2026-10-05', '2026-10-08', '', 6, 'PR #', 'Example row'),
]
FIRST, LAST = 5, 54
for i, row in enumerate(examples):
    for c, v in enumerate(row, 1):
        wk.cell(row=FIRST + i, column=c, value=v)
for r in range(FIRST, LAST + 1):
    for c in (7, 8, 9):
        wk.cell(row=r, column=c).number_format = 'yyyy-mm-dd'
    wk.cell(row=r, column=10).number_format = '0.0'
body(wk, FIRST, LAST, len(cols))
for r in range(FIRST, FIRST + len(examples)):
    for c in range(1, len(cols) + 1):
        wk.cell(row=r, column=c).font = Font(italic=True, color='71717A', name='Calibri')
for dv, col in ((dv_member_new(), 'D'), (dv_priority, 'E'), (dv_status, 'F')):
    wk.add_data_validation(dv)
    dv.add(f'{col}{FIRST}:{col}{LAST}')
wk.conditional_formatting.add(f'F{FIRST}:F{LAST}', CellIsRule(operator='equal', formula=['"Done"'], fill=PatternFill('solid', bgColor='BBF7D0', fgColor='BBF7D0')))
wk.conditional_formatting.add(f'F{FIRST}:F{LAST}', CellIsRule(operator='equal', formula=['"In progress"'], fill=PatternFill('solid', bgColor='FEF08A', fgColor='FEF08A')))
wk.conditional_formatting.add(f'F{FIRST}:F{LAST}', CellIsRule(operator='equal', formula=['"Blocked"'], fill=PatternFill('solid', bgColor='FECACA', fgColor='FECACA')))
wk.conditional_formatting.add(f'A{FIRST}:L{LAST}', FormulaRule(formula=[f'AND($H{FIRST}<>"",$F{FIRST}<>"Done",$H{FIRST}<TODAY())'], font=Font(color='B91C1C', bold=True)))
wk.freeze_panes = 'A5'
wk.auto_filter.ref = f'A4:L{LAST}'

# ------------------------------------------------------------------------------------------------ Weekly progress
wp = wb.create_sheet('Weekly Progress')
title(wp, 'Weekly Progress Report', 'Each member fills one row per week or sprint.')
pcols = ['Week / Sprint', 'Member', 'Planned tasks', 'Completed tasks', '% done', 'Hours', 'Issues / blockers', 'Next steps']
header(wp, 4, pcols, [14, 26, 40, 40, 9, 9, 36, 36])
PF, PL = 5, 64
_d = dv_member_new()
wp.add_data_validation(_d)
_d.add(f'B{PF}:B{PL}')
for r in range(PF, PL + 1):
    wp.cell(row=r, column=5).number_format = '0%'
    wp.cell(row=r, column=6).number_format = '0.0'
body(wp, PF, PL, len(pcols))
wp.freeze_panes = 'A5'
wp.auto_filter.ref = f'A4:H{PL}'

# ------------------------------------------------------------------------------------------------ Meeting log
wm = wb.create_sheet('Meeting Log')
title(wm, 'Meeting Log', 'Date, who came, what was decided, and who does what by when.')
mcols = ['Date', 'Type', 'Attendees', 'Agenda', 'Decisions', 'Action items', 'Owner', 'Due date']
header(wm, 4, mcols, [12, 16, 34, 36, 40, 40, 24, 12])
MF, ML = 5, 34
for r in range(MF, ML + 1):
    wm.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    wm.cell(row=r, column=8).number_format = 'yyyy-mm-dd'
body(wm, MF, ML, len(mcols))
dv_type = DataValidation(type='list', formula1='"Sprint planning,Daily check-in,Sprint review,Retrospective,Other"', allow_blank=True)
wm.add_data_validation(dv_type)
dv_type.add(f'B{MF}:B{ML}')
_d = dv_member_new()
wm.add_data_validation(_d)
_d.add(f'G{MF}:G{ML}')
wm.freeze_panes = 'A5'

# ------------------------------------------------------------------------------------------------ Peer evaluation
wv = wb.create_sheet('Peer Evaluation')
title(wv, 'Peer Evaluation', 'Rate every teammate from 1 (poor) to 5 (excellent). Do not rate yourself. Averages are calculated.')
crit = ['Contribution (amount of work)', 'Quality of work', 'Reliability (meets deadlines)', 'Communication', 'Teamwork and respect', 'Initiative / problem solving']
vcols = ['Evaluator', 'Evaluated member'] + crit + ['Average', 'Comments']
header(wv, 4, vcols, [26, 26] + [16] * 6 + [10, 44])
r = 5
for a in NAMES:
    for b in NAMES:
        if a != b:
            wv.cell(row=r, column=1, value=a)
            wv.cell(row=r, column=2, value=b)
            wv.cell(row=r, column=9, value=f'=IF(COUNT(C{r}:H{r})=0,"",ROUND(AVERAGE(C{r}:H{r}),2))')
            wv.cell(row=r, column=9).number_format = '0.00'
            wv.cell(row=r, column=9).fill = SUB_FILL
            r += 1
VL = r - 1
body(wv, 5, VL, len(vcols), banded=False)
for rr in range(5, VL + 1):
    wv.cell(row=rr, column=9).fill = SUB_FILL
    for cc in range(3, 9):
        wv.cell(row=rr, column=cc).alignment = CENTER
    wv.cell(row=rr, column=9).alignment = CENTER
wv.add_data_validation(dv_rating)
dv_rating.add(f'C5:H{VL}')
wv.freeze_panes = 'C5'
wv.auto_filter.ref = f'A4:J{VL}'

# ------------------------------------------------------------------------------------------------ Summary
wsu = wb.create_sheet('Summary')
title(wsu, 'Summary per member', 'Calculated from the other sheets. Do not type in the grey cells.')
scols = ['Member', 'Role', 'Tasks assigned', 'Tasks done', '% done', 'Hours (tasks)', 'Share of total hours', 'Average peer score (1 to 5)', 'Remark']
header(wsu, 4, scols, [28, 34, 14, 12, 10, 14, 16, 20, 40])
TL = LAST
for i in range(len(TEAM)):
    r = 5 + i
    wsu.cell(row=r, column=1, value=f'=Team!B{5 + i}')
    wsu.cell(row=r, column=2, value=f'=Team!C{5 + i}')
    wsu.cell(row=r, column=3, value=f"=COUNTIF('Task Tracker'!$D${FIRST}:$D${TL},A{r})")
    wsu.cell(row=r, column=4, value=f"=COUNTIFS('Task Tracker'!$D${FIRST}:$D${TL},A{r},'Task Tracker'!$F${FIRST}:$F${TL},\"Done\")")
    wsu.cell(row=r, column=5, value=f'=IF(C{r}=0,"",D{r}/C{r})')
    wsu.cell(row=r, column=6, value=f"=SUMIF('Task Tracker'!$D${FIRST}:$D${TL},A{r},'Task Tracker'!$J${FIRST}:$J${TL})")
    wsu.cell(row=r, column=7, value=f'=IF(SUM($F$5:$F${4 + len(TEAM)})=0,"",F{r}/SUM($F$5:$F${4 + len(TEAM)}))')
    wsu.cell(row=r, column=8, value=f"=IFERROR(ROUND(AVERAGEIF('Peer Evaluation'!$B$5:$B${VL},A{r},'Peer Evaluation'!$I$5:$I${VL}),2),\"\")")
    wsu.cell(row=r, column=5).number_format = '0%'
    wsu.cell(row=r, column=6).number_format = '0.0'
    wsu.cell(row=r, column=7).number_format = '0%'
    wsu.cell(row=r, column=8).number_format = '0.00'
    for c in range(1, 9):
        wsu.cell(row=r, column=c).fill = SUB_FILL
body(wsu, 5, 4 + len(TEAM), len(scols), banded=False)
for r in range(5, 5 + len(TEAM)):
    for c in range(1, 9):
        wsu.cell(row=r, column=c).fill = SUB_FILL
tr = 5 + len(TEAM)
wsu.cell(row=tr, column=1, value='Total')
for col in (3, 4, 6):
    L = get_column_letter(col)
    wsu.cell(row=tr, column=col, value=f'=SUM({L}5:{L}{tr - 1})')
wsu.cell(row=tr, column=5, value=f'=IF(C{tr}=0,"",D{tr}/C{tr})')
wsu.cell(row=tr, column=5).number_format = '0%'
wsu.cell(row=tr, column=6).number_format = '0.0'
for c in range(1, 10):
    wsu.cell(row=tr, column=c).font = Font(bold=True, name='Calibri')
    wsu.cell(row=tr, column=c).border = BORDER
wsu.freeze_panes = 'A5'

# order: Instructions, Team, Task Tracker, Weekly Progress, Meeting Log, Peer Evaluation, Summary
wb._sheets = [wb['Instructions'], wb['Team'], wb['Task Tracker'], wb['Weekly Progress'], wb['Meeting Log'], wb['Peer Evaluation'], wb['Summary']]
for sheet in wb.worksheets:
    sheet.sheet_view.zoomScale = 100
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print('wrote', OUT, OUT.stat().st_size, 'bytes')
